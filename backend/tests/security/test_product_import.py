from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from conftest import register, write
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.domain.errors import BusinessError
from app.infrastructure.management import ManagementRepository
from app.infrastructure.models import AuditLog, Product, Role, RolePermission, User, UserRole
from app.infrastructure.product_spreadsheet import HEADERS, MAX_BYTES, XLSX_TYPE, spreadsheet_template
from app.presentation.product_import import parse_products


def workbook(rows=None, headers=HEADERS):
    book = Workbook()
    book.active.append(headers)
    for row in rows or [product_row()]:
        book.active.append(row)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def product_row(pid='excel-phone-01'):
    return [pid, 'گوشی آزمایشی', 'Test phone', 'smartphones', 123456789, 7, 'Example', 'توضیحات', 'RAM: 8 GB;Storage: 256 GB', '']


def grant(engine, write_access=True):
    with Session(engine) as db, db.begin():
        user = db.scalar(select(User).where(User.email == 'user@example.com'))
        db.add(Role(name='importer'))
        db.flush()
        permissions = ['admin:access', 'products:read'] + (['products:write'] if write_access else [])
        db.add_all(RolePermission(role='importer', permission=p) for p in permissions)
        db.add(UserRole(user_id=user.id, role='importer'))


def upload(client, raw=None, commit=False, mime=XLSX_TYPE):
    token = client.get('/api/csrf').json()['token']
    return client.post('/api/admin/products/import', params={'commit': str(commit).lower()},
                       content=raw if raw is not None else workbook(),
                       headers={'X-CSRF-TOKEN': token, 'Content-Type': mime})


def test_import_requires_auth_permission_csrf_and_limits(client, engine):
    assert upload(client).status_code == 401
    register(client)
    assert upload(client).status_code == 403
    grant(engine)
    assert client.post('/api/admin/products/import', content=workbook(),
                       headers={'Content-Type': XLSX_TYPE}).status_code == 400
    assert upload(client, mime='application/json').status_code == 415
    assert upload(client, b'x' * (MAX_BYTES + 1)).status_code == 413
    assert upload(client, b'not an excel file').status_code == 400


def test_read_only_admin_cannot_import_or_download_template(client, engine):
    register(client)
    grant(engine, False)
    assert upload(client).status_code == 403
    assert client.get('/api/admin/products/import/template').status_code == 403


def test_template_preview_commit_hidden_until_photo_and_no_duplicate_retry(client, engine):
    register(client)
    grant(engine)
    template = client.get('/api/admin/products/import/template')
    assert template.status_code == 200
    assert template.headers['content-type'] == XLSX_TYPE
    assert parse_products(template.content)[1] == []
    preview = upload(client).json()
    assert preview['canImport'] and preview['created'] == 0
    with Session(engine) as db:
        assert db.get(Product, 'excel-phone-01') is None
    result = upload(client, commit=True)
    assert result.status_code == 200, result.text
    assert result.json()['created'] == 1
    assert client.get('/api/products/excel-phone-01').status_code == 404
    with Session(engine) as db:
        product = db.get(Product, 'excel-phone-01')
        assert not product.is_active
        assert product.stock == 7
        assert round(product.price * 600000) == 123456789
        assert product.payload['specs'] == ['RAM: 8 GB', 'Storage: 256 GB']
        assert db.scalar(select(AuditLog).where(AuditLog.subject_id == 'excel-phone-01')).event == 'PRODUCT_CREATED'
    retry = upload(client, commit=True).json()
    assert retry['created'] == 0 and retry['errors'][0]['row'] == 2
    rows, _ = parse_products(workbook())
    data = rows[0]['product']
    data.pop('price')
    data['expectedStock'] = 7
    data['isActive'] = True
    assert write(client, 'PUT', '/api/admin/products/excel-phone-01', json=data).status_code == 400
    data['image'] = '/assets/products/phone-01.webp'
    assert write(client, 'PUT', '/api/admin/products/excel-phone-01', json=data).status_code == 200
    assert client.get('/api/products/excel-phone-01').status_code == 200


def test_invalid_batch_does_not_save_valid_rows(client, engine):
    register(client)
    grant(engine)
    invalid = product_row('excel-phone-02')
    invalid[4] = -3
    result = upload(client, workbook([product_row(), invalid]), commit=True).json()
    assert result['created'] == 0 and not result['canImport']
    assert result['errors'][0]['row'] == 3
    with Session(engine) as db:
        assert db.get(Product, 'excel-phone-01') is None


def test_existing_id_blocks_entire_batch_even_when_archived(client, engine):
    register(client)
    grant(engine)
    with Session(engine) as db, db.begin():
        db.get(Product, 'phone-01').is_active = False
    result = upload(client, workbook([product_row(), product_row('phone-01')]), commit=True).json()
    assert result['created'] == 0 and result['errors'][0]['row'] == 3
    with Session(engine) as db:
        assert db.get(Product, 'excel-phone-01') is None


def test_import_rolls_back_products_and_audit_on_mid_batch_failure(client, engine, monkeypatch):
    register(client)
    grant(engine)
    original = ManagementRepository.save_product

    def fail_second(self, pid, data, create=False):
        if pid == 'excel-phone-02':
            raise BusinessError(409, 'Concurrent change')
        return original(self, pid, data, create)

    monkeypatch.setattr(ManagementRepository, 'save_product', fail_second)
    result = upload(client, workbook([product_row(), product_row('excel-phone-02')]), commit=True)
    assert result.status_code == 409
    with Session(engine) as db:
        assert db.get(Product, 'excel-phone-01') is None
        assert db.scalar(select(AuditLog).where(AuditLog.subject_id == 'excel-phone-01')) is None


def test_import_rate_limit(client, engine, app):
    register(client)
    grant(engine)
    app.state.settings.email_account_limit = 1
    assert upload(client).status_code == 200
    assert upload(client).status_code == 429


def test_persian_headers_digits_categories_and_optional_columns():
    row = product_row()
    row[3:6] = ['گوشی هوشمند', '۱۲۳٬۴۵۶٬۷۸۹', '۷']
    rows, errors = parse_products(workbook([row[:6]], HEADERS[:6]))
    assert errors == []
    assert rows[0]['product']['priceRial'] == 123456789
    assert rows[0]['product']['brand'] == ''
    assert rows[0]['product']['category'] == 'smartphones'


@pytest.mark.parametrize('index,value', [(0, 'Bad ID'), (3, 'unknown'), (4, '1.25'), (5, -1), (4, '=1+1'), (5, True)])
def test_bad_cells_have_row_errors(index, value):
    row = product_row()
    row[index] = value
    _, errors = parse_products(workbook([row]))
    assert errors and errors[0]['row'] == 2


def test_duplicate_ids_headers_row_limit_and_empty_workbook():
    assert parse_products(workbook([product_row(), product_row()]))[1][0]['row'] == 3
    for raw in [workbook(headers=['id', 'id']), workbook([product_row()] * 501), workbook([[]])]:
        with pytest.raises(BusinessError):
            parse_products(raw)


def test_rejects_zip_bomb_and_xml_entities():
    output = BytesIO()
    with ZipFile(output, 'w', ZIP_DEFLATED) as archive:
        archive.writestr('oversized.xml', b'x' * (12 * 1024 * 1024 + 1))
    with pytest.raises(BusinessError):
        parse_products(output.getvalue())
    output = BytesIO()
    with ZipFile(BytesIO(workbook())) as source, ZipFile(output, 'w', ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == 'xl/worksheets/sheet1.xml':
                content = b'<!DOCTYPE worksheet [<!ENTITY evil "expanded">]>' + content.replace(b'Test phone', b'&evil;')
            target.writestr(entry.filename, content)
    with pytest.raises(BusinessError):
        parse_products(output.getvalue())


def test_template_has_no_formula_and_all_categories():
    book = load_workbook(BytesIO(spreadsheet_template()))
    assert len(book['Guide']['A']) >= 18
    assert all(cell.data_type != 'f' for sheet in book for row in sheet for cell in row)
    book.close()
