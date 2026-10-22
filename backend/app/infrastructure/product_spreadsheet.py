"""Bounded XLSX input/output adapter. Never evaluate formulas or follow workbook links."""
from io import BytesIO
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.domain.errors import BusinessError

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 500
MAX_COLUMNS = 20
XLSX_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
HEADERS = ['شناسه', 'نام فارسی', 'نام انگلیسی', 'دسته‌بندی', 'قیمت ریال', 'موجودی', 'برند', 'توضیحات', 'مشخصات', 'برچسب']
CATEGORIES = {
    'smartphones': 'گوشی هوشمند', 'laptops': 'لپ‌تاپ و پی‌سی', 'audio': 'تجهیزات صوتی',
    'gaming': 'گیمینگ و کنسول', 'smart-home': 'خانه هوشمند', 'wearables': 'پوشیدنی‌ها',
    'accessories': 'لوازم جانبی', 'cameras': 'دوربین و لنز', 'drones': 'پهپاد و کوادکوپتر',
    'more': 'سایر دسته‌ها',
}


def read_spreadsheet(raw):
    if len(raw) > MAX_BYTES:
        raise BusinessError(413, 'حداکثر حجم فایل اکسل ۲ مگابایت است.')
    book = None
    try:
        with ZipFile(BytesIO(raw)) as archive:
            entries = archive.infolist()
            if (len(entries) > 200 or sum(e.file_size for e in entries) > 12 * 1024 * 1024
                    or any(e.flag_bits & 1 or 'vbaproject' in e.filename.lower() for e in entries)):
                raise BusinessError(400, 'فایل اکسل بیش از حد بزرگ، رمزدار یا دارای ماکرو است.')
        book = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        if not book.worksheets:
            raise BusinessError(400, 'فایل اکسل برگهٔ محصولات ندارد.')
        sheet = book.worksheets[0]
        if sheet.max_column and sheet.max_column > MAX_COLUMNS:
            raise BusinessError(400, 'فایل حداکثر ۲۰ ستون دارد؛ ستون‌های اضافی را حذف کنید.')
        # Ignore untrusted dimensions; bound actual rows and columns independently.
        sheet.reset_dimensions()
        rows = []
        for index, cells in enumerate(sheet.iter_rows(max_col=MAX_COLUMNS + 1), start=1):
            if index > MAX_ROWS + 1:
                raise BusinessError(400, 'هر فایل حداکثر ۵۰۰ ردیف محصول دارد؛ فایل را تقسیم کنید.')
            if cells[MAX_COLUMNS].value is not None:
                raise BusinessError(400, 'فایل حداکثر ۲۰ ستون دارد؛ از قالب نمونه استفاده کنید.')
            rows.append([(cell.value, cell.data_type) for cell in cells[:MAX_COLUMNS]])
        return rows
    except BusinessError:
        raise
    except Exception:
        # Workbook parsers can fail with XML, type, value and decompression errors.
        raise BusinessError(400, 'فایل معتبر XLSX انتخاب کنید؛ فایل رمزدار یا خراب قابل خواندن نیست.') from None
    finally:
        if book is not None:
            book.close()


def spreadsheet_template():
    book = Workbook()
    sheet = book.active
    sheet.title = 'Products'
    sheet.sheet_view.rightToLeft = True
    sheet.append(HEADERS)
    sheet.append(['my-product-01', 'محصول نمونه؛ این ردیف را جایگزین کنید', 'Example product',
                  'smartphones', 120000000, 5, 'Example', 'توضیحات محصول', 'رم: ۸ گیگابایت\nحافظه: ۲۵۶ گیگابایت', ''])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = 'A1:J2'
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='7C3AED')
    for i in range(1, 11):
        sheet.column_dimensions[get_column_letter(i)].width = 28
    guide = book.create_sheet('Guide')
    guide.sheet_view.rightToLeft = True
    guide.append(['راهنما', 'مقدار'])
    guide.append(['فقط برگهٔ اول وارد می‌شود', 'حداکثر ۵۰۰ محصول و حجم ۲ مگابایت؛ بدون فرمول'])
    guide.append(['ستون‌های الزامی', 'شناسه، نام فارسی، نام انگلیسی، دسته‌بندی، قیمت ریال، موجودی'])
    guide.append(['شناسه', 'یکتا؛ حروف کوچک انگلیسی، عدد و خط تیره؛ حداکثر ۸۰ کاراکتر'])
    guide.append(['قیمت و موجودی', 'عدد صحیح غیرمنفی؛ قیمت به ریال'])
    guide.append(['مشخصات', 'هر مشخصه در یک خط یا جداشده با ;'])
    guide.append(['وضعیت پس از ورود', 'بایگانی؛ عکس را در ویرایش اضافه و محصول را فعال کنید'])
    guide.append(['شناسه دسته‌بندی', 'عنوان'])
    for key, title in CATEGORIES.items():
        guide.append([key, title])
    guide.column_dimensions['A'].width = 35
    guide.column_dimensions['B'].width = 85
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()
